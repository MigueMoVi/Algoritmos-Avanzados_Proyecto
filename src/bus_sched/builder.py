"""Construcción de instancias.

Flujo general (igual para todas las rutas):
    datos de la ruta -> parámetros del modelo -> generación de vueltas
    -> ventanas de inicio -> instancia (el tráfico se aplica al resolver)

* window_policy_jobs: regla general de generación de vueltas y ventanas.
* from_excel: lee 'Parametros del modelo' (35 empresas/rutas), genera las
  vueltas con la regla general y verifica que coincidan con la hoja
  'Trabajos y ventanas' del Excel del proyecto.
* manual_instance: instancia pequeña del ejemplo resuelto a mano.
* random_instance: generador reproducible (semilla) para pruebas.
"""

from __future__ import annotations

import random
from pathlib import Path

from .model import Instance, Job, make_buses
from .timeutil import hhmm_to_min

DAY_START, DAY_END = 360, 1320          # 06:00 y 22:00 en minutos


def window_policy_jobs(route: str, n: int, base_min: float) -> list[Job]:
    """Regla general de generación de vueltas y ventanas.

    * última salida = round(22:00 - duración base);
    * n inicios nominales equiespaciados entre 06:00 y la última salida,
      redondeados al minuto (round de Python: mitad al par);
    * ventana ±15 min si el inicio nominal está en 06:00–09:00 o 16:00–19:00,
      ±30 min en el resto, recortada a [06:00, última salida].
    """
    latest = round(DAY_END - base_min)
    out = []
    for k in range(n):
        nominal = round(DAY_START + k * (latest - DAY_START) / max(1, n - 1))
        half = 15 if (nominal < 540 or 960 <= nominal < 1140) else 30
        out.append(Job(f"{route}-J{k + 1:03d}", route, base_min, max(DAY_START, nominal - half),
                       min(latest, nominal + half), nominal))
    return out


def check_instance(inst: Instance) -> list[str]:
    """Validación de datos de una instancia: ventanas válidas, vueltas que
    pueden terminar dentro de la jornada, duraciones positivas, ids únicos."""
    issues = []
    if len({j.id for j in inst.jobs}) != inst.n:
        issues.append(f"{inst.route}: identificadores de vuelta repetidos")
    for j in inst.jobs:
        if j.base <= 0:
            issues.append(f"{j.id}: duración base no positiva")
        if j.r > j.d:
            issues.append(f"{j.id}: ventana inválida (r > d)")
        # tolerancia de 1 min: las horas se expresan en HH:MM (redondeo al minuto)
        if j.r < inst.day_start or j.d + j.base > inst.day_end + 1.0:
            issues.append(f"{j.id}: la vuelta no cabe en la jornada")
    if inst.m <= 0:
        issues.append(f"{inst.route}: flota vacía")
    return issues


def from_excel(path: str | Path) -> list[Instance]:
    """Construye una instancia por ruta a partir del Excel del proyecto.

    Para cada ruta: flota y duración base de 'Parametros del modelo', vueltas
    generadas con window_policy_jobs y comparación con 'Trabajos y ventanas'.
    Cualquier discrepancia o dato inválido detiene la construcción.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    params = {}
    for row in wb["Parametros del modelo"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        params[row[0]] = {"company": row[1], "m": int(row[2]), "base_h": float(row[3]),
                          "demanda_h": row[4], "viajes_plan": int(row[6])}
    sheet: dict[str, list[tuple]] = {r: [] for r in params}
    for row in wb["Trabajos y ventanas"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        jid, route, _comp, _bus, _base_h, _per, nominal, r, d = row[:9]
        if route not in sheet:
            raise ValueError(f"'Trabajos y ventanas' contiene una ruta sin parámetros: {route}")
        sheet[route].append((jid, hhmm_to_min(nominal), hhmm_to_min(r), hhmm_to_min(d)))

    out, errors = [], []
    for route, p in params.items():
        jobs = window_policy_jobs(route, p["viajes_plan"], p["base_h"] * 60)
        got = sheet[route]
        if len(got) != len(jobs):
            errors.append(f"{route}: {len(got)} vueltas en la hoja, {len(jobs)} planificadas")
        else:
            for j, (jid, nom, r, d) in zip(jobs, got):
                if (j.id, j.nominal, j.r, j.d) != (jid, nom, r, d):
                    errors.append(f"{route}: la vuelta {jid} no coincide con la regla general")
                    break
        inst = Instance(route=route, company=p["company"], jobs=jobs,
                        buses=make_buses(route, p["m"]),
                        meta={"fuente": "Excel del proyecto (35 empresas/rutas)",
                              "duracion_base_h": p["base_h"],
                              "demanda_pasajeros_h": p["demanda_h"],
                              "viajes_planificados": p["viajes_plan"]})
        errors.extend(check_instance(inst))
        out.append(inst)
    if errors:
        raise ValueError("Datos inconsistentes: " + "; ".join(errors[:5]))
    return out


def manual_instance() -> Instance:
    """Instancia del ejemplo manual de la monografía.

    Ruta pequeña con duración base de 90 min (cercana a la de RTI-08, 1.57 h,
    redondeada para el cálculo a mano). 3 buses, 8 vueltas; almuerzo con
    ventanas 10:30–11:30, 12:00–13:00 y 13:30–14:30.
    """
    route = "DEMO"
    spec = [  # id, r, d  (política de ventanas del proyecto)
        ("J1", "06:00", "06:30"),
        ("J2", "06:30", "07:00"),
        ("J3", "07:00", "07:30"),
        ("J4", "08:00", "08:30"),
        ("J5", "09:30", "10:30"),
        ("J6", "11:30", "12:30"),
        ("J7", "16:30", "17:00"),
        ("J8", "17:00", "17:30"),
    ]
    jobs = [Job(j, route, 90.0, hhmm_to_min(r), hhmm_to_min(d),
                (hhmm_to_min(r) + hhmm_to_min(d)) / 2) for j, r, d in spec]
    return Instance(route=route, company="Ruta del ejemplo manual", jobs=jobs,
                    buses=make_buses(route, 3),
                    meta={"nota": "Instancia didáctica para el ejemplo manual"})


def random_instance(n: int, m: int, seed: int = 0, base: float = 120.0,
                    spread: float = 0.0, route: str = "RND") -> Instance:
    """n vueltas con inicio nominal uniforme en la jornada y ventana según la política del proyecto."""
    rng = random.Random(seed)
    jobs = []
    latest = 22 * 60 - base * 1.5
    for k in range(n):
        nominal = rng.uniform(360, latest)
        half = 15 if (360 <= nominal < 540 or 960 <= nominal < 1140) else 30
        b = base * (1 + rng.uniform(-spread, spread))
        r, d = max(360.0, nominal - half), min(latest, nominal + half)
        jobs.append(Job(f"{route}-J{k + 1:03d}", route, round(b, 2), round(r), round(d),
                        round(nominal)))
    return Instance(route=route, company="Sintética", jobs=jobs,
                    buses=make_buses(route, m), meta={"seed": seed})
