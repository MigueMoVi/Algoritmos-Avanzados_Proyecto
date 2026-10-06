"""Construcción de instancias.

* from_excel: lee la tabla del proyecto (37 empresas/rutas) y genera una
  instancia JSON por ruta (hojas 'Parametros del modelo' y
  'Trabajos y ventanas').
* manual_instance: instancia didáctica del ejemplo resuelto a mano.
* random_instance: generador reproducible (semilla) para pruebas y escala.
"""

from __future__ import annotations

import random
from pathlib import Path

from .model import Instance, Job, make_buses
from .timeutil import hhmm_to_min


def option_c_jobs(route: str, n: int, base_min: float) -> list[Job]:
    """Regla de la Opción C (reconstruida de la hoja 'Trabajos y ventanas'):
    n inicios nominales equiespaciados entre 06:00 y la última salida que
    termina a las 22:00; ventana ±15 min en 06–09 y 16–19, ±30 min en el
    resto, recortada a [06:00, última salida]."""
    import math
    latest = math.floor(22 * 60 - base_min)
    out = []
    for k in range(n):
        nominal = math.floor(360 + k * (latest - 360) / max(1, n - 1) + 0.5)
        half = 15 if (nominal < 540 or 960 <= nominal < 1140) else 30
        out.append(Job(f"{route}-J{k + 1:03d}", route, base_min, max(360, nominal - half),
                       min(latest, nominal + half), nominal))
    return out


def from_excel(path: str | Path) -> list[Instance]:
    """Lee el Excel. Control de calidad: si el número de vueltas de la hoja
    'Trabajos y ventanas' no coincide con 'Viajes planificados de la flota'
    (hoja 'Parametros del modelo'), la ruta se regenera con option_c_jobs y
    se marca en meta['correccion'] (caso RTU-28 y RTU-29: columnas corridas
    en la tabla fuente, 98.6 y 94.52 "viajes por unidad")."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    params = {}
    for row in wb["Parametros del modelo"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        params[row[0]] = {"company": row[1], "m": int(row[2]), "base_h": float(row[3]),
                          "demanda_h": row[4], "viajes_plan": row[6]}
    jobs: dict[str, list[Job]] = {r: [] for r in params}
    for row in wb["Trabajos y ventanas"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        jid, route, _comp, _bus, base_h, _per, nominal, r, d = row[:9]
        jobs[route].append(Job(jid, route, float(base_h) * 60, hhmm_to_min(r),
                               hhmm_to_min(d), hhmm_to_min(nominal)))
    out = []
    for route, p in params.items():
        meta = {"fuente": "Tabla 37 empresas (Opción C)", "duracion_base_h": p["base_h"],
                "demanda_pasajeros_h": p["demanda_h"], "viajes_planificados": p["viajes_plan"]}
        js = jobs[route]
        if p["viajes_plan"] and len(js) != int(p["viajes_plan"]):
            meta["correccion"] = (f"La hoja de trabajos tenía {len(js)} vueltas; se regeneran "
                                  f"{int(p['viajes_plan'])} según 'Parametros del modelo'")
            js = option_c_jobs(route, int(p["viajes_plan"]), p["base_h"] * 60)
        out.append(Instance(route=route, company=p["company"], jobs=js,
                            buses=make_buses(route, p["m"]), meta=meta))
    return out


def manual_instance() -> Instance:
    """Instancia del ejemplo manual de la monografía.

    Ruta didáctica derivada de RTI-08 (CRISTO BLANCO S.A., tiempo de vuelta
    1.57 h); la duración base se redondea a 90 min para el cálculo a mano.
    3 buses, 8 vueltas, almuerzo en turnos 10:00 / 12:00 / 14:00.
    """
    route = "DEMO"
    spec = [  # id, r, d  (ventanas según Opción C)
        ("J1", "06:00", "06:30"),
        ("J2", "06:30", "07:00"),
        ("J3", "07:00", "07:30"),
        ("J4", "08:00", "08:30"),
        ("J5", "09:30", "10:30"),
        ("J6", "11:30", "12:30"),
        ("J7", "16:30", "17:00"),
        ("J8", "17:00", "17:30"),
    ]
    jobs = [Job(j, route, 90.0, hhmm_to_min(r), hhmm_to_min(d),
                (hhmm_to_min(r) + hhmm_to_min(d)) / 2) for j, r, d in spec]
    return Instance(route=route, company="Ruta didáctica (base RTI-08)", jobs=jobs,
                    buses=make_buses(route, 3),
                    meta={"nota": "Instancia didáctica para el ejemplo manual"})


def random_instance(n: int, m: int, seed: int = 0, base: float = 120.0,
                    spread: float = 0.0, route: str = "RND") -> Instance:
    """n vueltas con inicio nominal uniforme en la jornada y ventana Opción C."""
    rng = random.Random(seed)
    jobs = []
    latest = 22 * 60 - base * 1.5
    for k in range(n):
        nominal = rng.uniform(360, latest)
        half = 15 if (360 <= nominal < 540 or 960 <= nominal < 1140) else 30
        b = base * (1 + rng.uniform(-spread, spread))
        r, d = max(360.0, nominal - half), min(latest, nominal + half)
        jobs.append(Job(f"{route}-J{k + 1:03d}", route, round(b, 2), round(r), round(d),
                        round(nominal)))
    return Instance(route=route, company="Sintética", jobs=jobs,
                    buses=make_buses(route, m), meta={"seed": seed})
