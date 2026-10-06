"""Genera el Excel actualizado de la Entrega 1 a partir del Excel original,
las instancias corregidas y los resultados de run_experiments.py.

    python experiments/build_excel.py
Salida: results/Tabla_37_Empresas_Entrega1_Trafico.xlsx
"""

from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from sched.algorithms import branch_and_bound, list_scheduling, lpt  # noqa: E402
from sched.algorithms.cpsat import cpsat_solve  # noqa: E402
from sched.builder import from_excel, manual_instance  # noqa: E402
from sched.timeutil import min_to_hhmm  # noqa: E402
from sched.traffic import RAW_BANDS, get_profile  # noqa: E402

SRC = ROOT / "data" / "raw" / "tabla_37_empresas.xlsx"
OUT = ROOT / "results" / "Tabla_37_Empresas_Entrega1_Trafico.xlsx"
RES = ROOT / "results"

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(name="Calibri", bold=True, color="FFFFFF")
SUB_FILL = PatternFill("solid", fgColor="D9EAF7")
INPUT_FONT = Font(name="Calibri", color="0000FF")
LINK_FONT = Font(name="Calibri", color="008000")
BOLD = Font(name="Calibri", bold=True)
TITLE = Font(name="Calibri", bold=True, size=13, color="1F4E78")
NOTE = Font(name="Calibri", italic=True, size=9, color="595959")
FLAG = PatternFill("solid", fgColor="FFF2CC")
KEY = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def header(ws, row, cols, start_col=1):
    for k, text in enumerate(cols):
        c = ws.cell(row=row, column=start_col + k, value=text)
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def widths(ws, ws_widths):
    for col, w in ws_widths.items():
        ws.column_dimensions[col].width = w


def t(hhmm: str) -> float:
    h, m = hhmm.split(":")
    return (int(h) * 60 + int(m)) / 1440


# ====================================================================
def sheet_trabajos(wb, insts):
    """Reescribe 'Trabajos y ventanas' con RTU-28 / RTU-29 corregidas."""
    ws = wb["Trabajos y ventanas"]
    hdr = [c.value for c in ws[1]]
    styles = [copy(c._style) for c in ws[2]]
    orig = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            orig.setdefault(row[1], []).append(row)
    ws.delete_rows(2, ws.max_row)
    r = 2
    by = {i.route: i for i in insts}
    for route in orig:
        inst = by[route]
        if "correccion" in inst.meta:
            for j in inst.jobs:
                nom = j.nominal
                per = ("06:00–09:00" if nom < 540 else "09:00–16:00" if nom < 960
                       else "16:00–19:00" if nom < 1140 else "19:00–22:00")
                vals = [j.id, route, inst.company, f"{route}-B1...B{inst.m}",
                        round(j.base / 60, 2), per, min_to_hhmm(nom), min_to_hhmm(j.r),
                        min_to_hhmm(j.d), int(round(j.d - j.r)),
                        "06:00–22:00 (2 h almuerzo escalonado)",
                        "CORREGIDO Entrega 1: regenerado con la regla Opción C a partir de "
                        "'Viajes planificados' (la versión anterior usaba un valor desplazado de la fuente)"]
                for k, v in enumerate(vals):
                    c = ws.cell(row=r, column=k + 1, value=v)
                    c._style = copy(styles[k])
                    if k == 11:
                        c.fill = FLAG
                r += 1
        else:
            for row in orig[route]:
                for k, v in enumerate(row):
                    ws.cell(row=r, column=k + 1, value=v)._style = copy(styles[k])
                r += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{r - 1}"


def flag_source(wb):
    ws = wb["Tabla proyecto"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value in ("RTU-28", "RTU-29"):
            for idx in (9, 10, 11):          # demanda/h, viajes por unidad, frecuencia requerida
                row[idx].fill = FLAG
            row[10].comment = Comment(
                "Valor anómalo (Entrega 1): las columnas Demanda_pasajeros_h, "
                "Viajes_por_unidad_vehicular y Frecuencia_requerida aparecen desplazadas "
                "respecto de la fuente. No se usa: el modelo toma 'Viajes planificados' de la "
                "hoja 'Parametros del modelo' (145 y 139).", "Grupo 4")


# ====================================================================
def sheet_perfil(wb):
    ws = wb.create_sheet("Perfil trafico")
    ws["A1"] = "MODELO DE TRÁFICO – Factores de congestión por franja horaria"
    ws["A1"].font = TITLE
    ws["A2"] = ("Factor = multiplicador del tiempo de recorrido dentro de la franja (velocidad "
                "constante por tramos, modelo de Ichoua–Gendreau–Potvin). Celdas azules = "
                "supuestos editables; T1 se recalcula sola.")
    ws["A2"].font = NOTE
    header(ws, 4, ["Inicio franja", "Fin franja", "Minutos", "Condición",
                   "Factor T2 (pesimista)", "Factor T1 (normalizado)", "Factor T0 (sin tráfico)"])
    for k, (a, b, f, lab) in enumerate(RAW_BANDS):
        r = 5 + k
        ws.cell(r, 1, t(a)).number_format = "hh:mm"
        ws.cell(r, 2, t(b)).number_format = "hh:mm"
        ws.cell(r, 3, f"=ROUND((B{r}-A{r})*1440,0)")
        ws.cell(r, 4, lab)
        c = ws.cell(r, 5, f)
        c.font = INPUT_FONT
        ws.cell(r, 6, f"=E{r}/$B$15").number_format = "0.000"
        ws.cell(r, 7, 1)
        for col in range(1, 8):
            ws.cell(r, col).border = BOX
    last = 5 + len(RAW_BANDS) - 1
    ws["A13"] = "Factor fuera de franjas (después de 22:00)"
    ws["B13"] = 0.9
    ws["B13"].font = INPUT_FONT
    ws["A14"] = "Minutos de jornada"
    ws["B14"] = f"=SUM(C5:C{last})"
    ws["A15"] = "Media ponderada del factor T2 (divisor de normalización)"
    ws["B15"] = f"=SUMPRODUCT(C5:C{last},E5:E{last})/B14"
    ws["B15"].number_format = "0.0000"
    ws["B15"].fill = KEY
    ws["A16"] = "Comprobación: media ponderada T1 (debe ser 1)"
    ws["B16"] = f"=SUMPRODUCT(C5:C{last},F5:F{last})/B14"
    ws["B16"].number_format = "0.0000"
    ws["A17"] = "Factor fuera de franjas T1"
    ws["B17"] = "=B13/B15"
    ws["B17"].number_format = "0.000"
    for r in range(13, 18):
        ws.cell(r, 1).font = BOLD
    ws["A19"] = ("Nota: los factores son SUPUESTOS EXPERIMENTALES (no hay un índice público de "
                 "tráfico para Cusco). T1 conserva la duración media de la tabla 2020; T2 es un "
                 "escenario de estrés. Fuente del modelo: Ichoua, Gendreau y Potvin (2003), EJOR 144(2).")
    ws["A19"].font = NOTE
    widths(ws, {"A": 16, "B": 12, "C": 10, "D": 32, "E": 18, "F": 20, "G": 18})
    ws.freeze_panes = "A5"


def sheet_calculadora(wb):
    ws = wb.create_sheet("Calculadora vuelta")
    ws["A1"] = "CALCULADORA – Duración de una vuelta según la hora de salida"
    ws["A1"].font = TITLE
    ws["A2"] = "Edite las celdas azules. La vuelta avanza franja por franja (propiedad FIFO)."
    ws["A2"].font = NOTE
    ws["A4"], ws["B4"] = "Hora de salida", t("06:00")
    ws["B4"].number_format = "hh:mm"
    ws["A5"], ws["B5"] = "Duración base (min)", 90
    ws["A6"], ws["B6"] = "Escenario (T0 / T1 / T2)", "T2"
    for c in ("B4", "B5", "B6"):
        ws[c].font = INPUT_FONT
        ws[c].fill = KEY
    ws["A7"], ws["B7"] = "Salida (min desde 00:00)", "=ROUND(B4*1440,4)"
    for r in range(4, 8):
        ws.cell(r, 1).font = BOLD
    header(ws, 9, ["Inicio franja (min)", "Fin franja (min)", "Factor escenario",
                   "Minutos reales disponibles", "Capacidad en min base", "Base acumulada previa",
                   "Base recorrida en franja", "Minutos reales en franja"])
    n = len(RAW_BANDS)
    for k in range(n + 1):
        r = 10 + k
        if k < n:
            src = 5 + k
            ws.cell(r, 1, f"=ROUND('Perfil trafico'!A{src}*1440,0)").font = LINK_FONT
            ws.cell(r, 2, f"=ROUND('Perfil trafico'!B{src}*1440,0)").font = LINK_FONT
            ws.cell(r, 3, f"=IF($B$6=\"T0\",1,IF($B$6=\"T1\",'Perfil trafico'!F{src},"
                          f"'Perfil trafico'!E{src}))").font = LINK_FONT
        else:   # tramo posterior a las 22:00
            ws.cell(r, 1, 1320)
            ws.cell(r, 2, 2880)
            ws.cell(r, 3, "=IF($B$6=\"T0\",1,IF($B$6=\"T1\",'Perfil trafico'!B17,"
                          "'Perfil trafico'!B13))").font = LINK_FONT
        ws.cell(r, 4, f"=MAX(0,B{r}-MAX(A{r},$B$7))")
        ws.cell(r, 5, f"=D{r}/C{r}")
        ws.cell(r, 6, "=0" if k == 0 else f"=F{r - 1}+G{r - 1}")
        ws.cell(r, 7, f"=MAX(0,MIN(E{r},$B$5-F{r}))")
        ws.cell(r, 8, f"=G{r}*C{r}")
        for col in range(1, 9):
            ws.cell(r, col).border = BOX
            if col >= 3:
                ws.cell(r, col).number_format = "0.000"
    lr = 10 + n
    ws[f"A{lr + 2}"], ws[f"B{lr + 2}"] = "Duración real (min)", f"=SUM(H10:H{lr})"
    ws[f"A{lr + 3}"], ws[f"B{lr + 3}"] = "Hora de llegada", f"=(B7+B{lr + 2})/1440"
    ws[f"B{lr + 3}"].number_format = "hh:mm"
    ws[f"A{lr + 4}"], ws[f"B{lr + 4}"] = "Sobrecosto por tráfico", f"=B{lr + 2}/B5-1"
    ws[f"B{lr + 4}"].number_format = "0.0%"
    ws[f"A{lr + 6}"] = "Verificación con el programa (salida 06:00, base 90, T2):"
    p = get_profile("T2").travel_time(360, 90)
    ws[f"B{lr + 6}"] = round(p, 4)
    ws[f"C{lr + 6}"] = "← valor de sched.traffic.travel_time(360, 90); debe coincidir con B" + str(lr + 2)
    ws[f"C{lr + 6}"].font = NOTE
    for r in (lr + 2, lr + 3, lr + 4, lr + 6):
        ws.cell(r, 1).font = BOLD
    ws[f"B{lr + 2}"].number_format = "0.00"
    ws[f"B{lr + 2}"].fill = KEY
    widths(ws, {"A": 26, "B": 14, "C": 14, "D": 16, "E": 16, "F": 16, "G": 16, "H": 16})


def sheet_saturacion(wb, n_routes):
    ws = wb.create_sheet("Saturacion flota")
    ws["A1"] = "SATURACIÓN DE LA FLOTA – ¿alcanza la flota para las vueltas planificadas?"
    ws["A1"].font = TITLE
    ws["A2"] = ("ρ = horas de vuelta requeridas / horas netas disponibles (14 h por bus). "
                "ρ > 1 implica infactibilidad segura aun sin tráfico ni ventanas.")
    ws["A2"].font = NOTE
    header(ws, 4, ["Ruta", "Empresa", "Buses", "Duración base (h)", "Vueltas planificadas",
                   "Horas requeridas", "Horas netas/bus", "Horas disponibles", "ρ (T0)",
                   "ρ con T2 (aprox.)", "Estado"])
    pm = "'Parametros del modelo'"
    for k in range(n_routes):
        r, s = 5 + k, 2 + k
        ws.cell(r, 1, f"={pm}!A{s}").font = LINK_FONT
        ws.cell(r, 2, f"={pm}!B{s}").font = LINK_FONT
        ws.cell(r, 3, f"={pm}!C{s}").font = LINK_FONT
        ws.cell(r, 4, f"={pm}!D{s}").font = LINK_FONT
        ws.cell(r, 5, f"={pm}!G{s}").font = LINK_FONT
        ws.cell(r, 6, f"=D{r}*E{r}")
        ws.cell(r, 7, f"={pm}!L{s}").font = LINK_FONT
        ws.cell(r, 8, f"=C{r}*G{r}")
        ws.cell(r, 9, f"=F{r}/H{r}").number_format = "0.000"
        ws.cell(r, 10, f"=I{r}*'Perfil trafico'!$B$15").number_format = "0.000"
        ws.cell(r, 11, f"=IF(I{r}>1,\"Infactible (capacidad)\",IF(J{r}>0.85,\"Crítica con tráfico\","
                       f"IF(I{r}>0.75,\"Ajustada\",\"Holgada\")))")
        for col in range(1, 12):
            ws.cell(r, col).border = BOX
        ws.cell(r, 6).number_format = "0.0"
        ws.cell(r, 8).number_format = "0.0"
    e = 4 + n_routes
    ws.cell(e + 2, 1, "Totales").font = BOLD
    ws.cell(e + 2, 5, f"=SUM(E5:E{e})")
    ws.cell(e + 2, 6, f"=SUM(F5:F{e})").number_format = "0.0"
    ws.cell(e + 2, 8, f"=SUM(H5:H{e})").number_format = "0.0"
    ws.cell(e + 2, 9, f"=F{e + 2}/H{e + 2}").number_format = "0.000"
    ws.cell(e + 3, 1, "Rutas con ρ > 1").font = BOLD
    ws.cell(e + 3, 2, f"=COUNTIF(I5:I{e},\">1\")")
    ws.cell(e + 4, 1, "Rutas críticas con tráfico T2 (ρ·factor > 0.85)").font = BOLD
    ws.cell(e + 4, 2, f"=COUNTIF(J5:J{e},\">0.85\")")
    widths(ws, {"A": 10, "B": 40, "C": 8, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12,
                "I": 10, "J": 12, "K": 22})
    ws.freeze_panes = "A5"


def sheet_manual(wb):
    ws = wb.create_sheet("Instancia manual")
    inst = manual_instance()
    t2 = get_profile("T2")
    ws["A1"] = "INSTANCIA MANUAL – 3 buses, 8 vueltas, base 90 min, tráfico T2"
    ws["A1"].font = TITLE
    ws["A2"] = ("Ruta didáctica derivada de RTI-08 (tiempo de vuelta redondeado a 90 min). "
                "Almuerzo 2 h: B1 inicia 10:30–11:30, B2 12:00–13:00, B3 13:30–14:30.")
    ws["A2"].font = NOTE
    header(ws, 4, ["Vuelta", "Inicio mín.", "Inicio máx.", "Duración si sale en inicio mín. (T2)",
                   "Duración si sale en inicio máx. (T2)"])
    for k, j in enumerate(inst.jobs):
        r = 5 + k
        ws.cell(r, 1, j.id)
        ws.cell(r, 2, j.r / 1440).number_format = "hh:mm"
        ws.cell(r, 3, j.d / 1440).number_format = "hh:mm"
        ws.cell(r, 4, round(t2.travel_time(j.r, j.base), 2))
        ws.cell(r, 5, round(t2.travel_time(j.d, j.base), 2))
        for col in range(1, 6):
            ws.cell(r, col).border = BOX
    r0 = 15
    ws.cell(r0 - 1, 1, "Asignaciones producidas por el programa (python -m sched demo --traffic T2)").font = BOLD
    header(ws, r0, ["Algoritmo", "Vuelta", "Bus", "Salida", "Llegada", "Duración (min)"])
    sols = [list_scheduling(inst, t2), lpt(inst, t2), branch_and_bound(inst, t2),
            cpsat_solve(inst, t2, time_limit_s=20)]
    r = r0 + 1
    for s in sols:
        for a in sorted(s.assignments.values(), key=lambda a: (a.bus_id, a.start)):
            ws.cell(r, 1, s.algorithm)
            ws.cell(r, 2, a.job_id)
            ws.cell(r, 3, a.bus_id.split("-")[-1])
            ws.cell(r, 4, a.start / 1440).number_format = "hh:mm"
            ws.cell(r, 5, a.end / 1440).number_format = "hh:mm"
            ws.cell(r, 6, round(a.duration, 2))
            for col in range(1, 7):
                ws.cell(r, col).border = BOX
            r += 1
    last = r - 1
    rr = r0
    header(ws, rr, ["Algoritmo", "Carga B1", "Carga B2", "Carga B3", "L_max (min)", "L_max (h)",
                    "Desbalance (min)"], start_col=9)
    for k, s in enumerate(sols):
        row = rr + 1 + k
        ws.cell(row, 9, s.algorithm)
        for b in range(3):
            col = 10 + b
            ws.cell(row, col, f"=SUMIFS($F${r0 + 1}:$F${last},$A${r0 + 1}:$A${last},$I{row},"
                              f"$C${r0 + 1}:$C${last},\"B{b + 1}\")").number_format = "0.0"
        ws.cell(row, 13, f"=MAX(J{row}:L{row})").number_format = "0.0"
        ws.cell(row, 14, f"=M{row}/60").number_format = "0.00"
        ws.cell(row, 15, f"=MAX(J{row}:L{row})-MIN(J{row}:L{row})").number_format = "0.0"
        ws.cell(row, 13).fill = KEY
        for col in range(9, 16):
            ws.cell(row, col).border = BOX
    widths(ws, {"A": 11, "B": 11, "C": 11, "D": 16, "E": 16, "F": 14, "I": 12, "J": 10,
                "K": 10, "L": 10, "M": 12, "N": 10, "O": 14})


def sheet_resultados(wb):
    df = pd.read_csv(RES / "E1_37_rutas.csv")
    d = df[df.politica == "earliest"]
    ws = wb.create_sheet("Resultados por ruta")
    ws["A1"] = "RESULTADOS PRELIMINARES POR RUTA (E1) – salida de experiments/run_experiments.py"
    ws["A1"].font = TITLE
    ws["A2"] = "Valores generados por el programa (no se editan a mano). L_max en horas; NA = vueltas no asignadas."
    ws["A2"].font = NOTE
    cols = ["Ruta", "Buses", "Vueltas", "ρ"]
    for tr in ("T0", "T1", "T2"):
        cols += [f"LS NA {tr}", f"LS L_max {tr}", f"LPT NA {tr}", f"LPT L_max {tr}"]
    header(ws, 4, cols)
    for k, route in enumerate(sorted(d.ruta.unique())):
        r = 5 + k
        base = d[(d.ruta == route)].iloc[0]
        vals = [route, int(base.buses), int(base.trabajos), round(float(base.rho_saturacion), 3)]
        for tr in ("T0", "T1", "T2"):
            for alg in ("LS", "LPT"):
                x = d[(d.ruta == route) & (d.trafico == tr) & (d.algoritmo == alg)].iloc[0]
                vals += [int(x.no_asignados), round(float(x.Lmax_h), 3)]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(r, c, v)
            cell.border = BOX
            if "NA" in cols[c - 1] and isinstance(v, int) and v > 0:
                cell.fill = FLAG
    e = 4 + d.ruta.nunique()
    ws.cell(e + 1, 1, "Total / media").font = BOLD
    for c in range(2, len(cols) + 1):
        L = get_column_letter(c)
        f = "SUM" if ("NA" in cols[c - 1] or cols[c - 1] in ("Buses", "Vueltas")) else "AVERAGE"
        ws.cell(e + 1, c, f"={f}({L}5:{L}{e})").font = BOLD
        if f == "AVERAGE":
            ws.cell(e + 1, c).number_format = "0.000"
    ws.freeze_panes = "B5"
    widths(ws, {get_column_letter(c): 10 for c in range(1, len(cols) + 1)})


def sheet_resumen(wb):
    ws = wb.create_sheet("Resumen experimentos")
    ws["A1"] = "RESUMEN DE EXPERIMENTOS PRELIMINARES – Entrega 1"
    ws["A1"].font = TITLE
    r = 3
    blocks = [
        ("E1 · 37 rutas × algoritmo × escenario de tráfico × política de inicio", "E1_resumen.csv"),
        ("E2 · Heurísticas vs exactos (B&B y CP-SAT), 45 instancias pequeñas por escenario",
         "E2_resumen.csv"),
        ("E3 · Escalabilidad (instancias sintéticas, m = n/3, perfil T1)", "E3_escalabilidad.csv"),
        ("E4 · Robustez Monte Carlo (T1, 100 réplicas por ruta)", "E4_resumen.csv"),
    ]
    for title, f in blocks:
        df = pd.read_csv(RES / f)
        ws.cell(r, 1, title).font = BOLD
        r += 1
        header(ws, r, list(df.columns))
        for row in df.itertuples(index=False):
            r += 1
            for c, v in enumerate(row, start=1):
                if isinstance(v, float):
                    v = round(v, 3)
                ws.cell(r, c, v).border = BOX
        r += 3
    widths(ws, {get_column_letter(c): 15 for c in range(1, 15)})


def notas(wb):
    ws = wb["Notas"]
    r = ws.max_row + 2
    rows = [
        ("ENTREGA 1 – Cambios", ""),
        ("Tráfico (observación del docente)",
         "La duración de cada vuelta depende de su hora de salida: p_j(s) se calcula integrando una "
         "velocidad constante por franja (hojas 'Perfil trafico' y 'Calculadora vuelta')."),
        ("Escenarios", "T0 sin tráfico (Avance 1), T1 normalizado (media 1), T2 pesimista. "
                       "Además, variabilidad estocástica LogNormal (σ = 0.10 y 0.20) en simulación."),
        ("Almuerzo", "2 h por bus con ventana de inicio de ±30 min en 3 turnos (11:00, 12:30, 14:00)."),
        ("Corrección de datos", "RTU-28 y RTU-29: 'Tabla proyecto' tiene columnas desplazadas (98.6 y "
                                "94.52 viajes/unidad). La hoja de trabajos se regeneró con 145 y 139 vueltas."),
        ("Objetivo", "Lexicográfico: 1) maximizar vueltas atendidas; 2) minimizar la carga máxima por bus L_max."),
        ("Saturación", "Hoja 'Saturacion flota': ρ > 1 indica que la flota no alcanza aunque no haya tráfico."),
    ]
    for a, b in rows:
        ws.cell(r, 1, a).font = BOLD
        ws.cell(r, 2, b).alignment = Alignment(wrap_text=True)
        r += 1


def main():
    insts = from_excel(SRC)
    wb = openpyxl.load_workbook(SRC)
    sheet_trabajos(wb, insts)
    flag_source(wb)
    sheet_perfil(wb)
    sheet_calculadora(wb)
    sheet_saturacion(wb, len(insts))
    sheet_manual(wb)
    sheet_resultados(wb)
    sheet_resumen(wb)
    notas(wb)
    wb.save(OUT)
    print("->", OUT)


if __name__ == "__main__":
    main()
